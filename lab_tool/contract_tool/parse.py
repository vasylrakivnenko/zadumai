"""Contract tool v0, document parsing (2026-10-05; the user: "we solve it for us"). Every format LAB hands an agent
(.docx, .xlsx, .pptx, .eml, .txt / .md, .pdf) becomes a list of units (a paragraph, a table row, a slide paragraph, an
email header block), each with its document, its location (the headings above it, its own clause number) and, for .docx,
the tracked changes kept apart: `text` = the current text (insertions in, deletions out: what LAB's `read` shows),
`before` = the original text, `marked` = "[-deleted-]{+inserted+}" (what `read` hides). General purpose: it knows nothing
about LAB's tasks or rubrics.
usage: dspy_venv/bin/python -m contract_tool.parse FILE   (prints the units)"""
import datetime, email, email.policy, html, os, re, shutil, subprocess, sys, zipfile
from dataclasses import dataclass, field
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
LABEL = re.compile(r"^\s*((?:ARTICLE|Article|SECTION|Section|SCHEDULE|Schedule|EXHIBIT|Exhibit|ANNEX|Annex|APPENDIX|Appendix|"
                   r"CLAUSE|Clause|PART|Part)\s+[\dIVXLC]+[A-Za-z]?(?:\.\d+)*|\d+(?:\.\d+)+|\d+\.(?=\s)|\([a-zA-Z0-9]{1,4}\)|[A-Z]\.(?=\s)|"
                   r"[ivx]{1,4}\.(?=\s))")


@dataclass
class Unit:
    doc: str
    i: int
    kind: str              # p (paragraph), row (table row), head (email headers), slide
    text: str              # current text
    before: str = ""       # original text (differs from text only where tracked changes exist)
    marked: str = ""       # text with [-deleted-] {+inserted+}
    changes: list = field(default_factory=list)  # (ins | del, text, author, date)
    heading: bool = False
    label: str = ""        # its own clause number: 8.2, (a), Article 5, row 7
    path: tuple = ()       # the headings above it, outermost first

    def where(self):
        p = " > ".join(h[:60] for h in self.path)
        return " | ".join(x for x in (self.doc, p, ("§" + self.label) if self.label and self.kind == "p" else self.label) if x)


def _clean(s):
    return re.sub(r"\s+\)", ")", re.sub(r"\(\s+", "(", re.sub(r"\s+", " ", s))).strip()


def label_of(text):
    m = LABEL.match(text)
    return m.group(1).rstrip(".") if m else ""


def _level(text):
    """1 = an article-level heading (ARTICLE / all caps / "8." / "8 Title"), 2 = a sub-heading."""
    if re.match(r"^(ARTICLE|Article|PART|Part|SCHEDULE|Schedule|EXHIBIT|Exhibit|ANNEX|Annex|APPENDIX|Appendix)\b", text) or text.isupper():
        return 1
    return 2 if re.match(r"^\d+\.\d+", text) else 1


class _Path:
    def __init__(self):
        self.p, self.num = (), ""

    def see(self, text):
        h = _clean(text)[:100]
        self.p = (h,) if _level(h) == 1 or not self.p else (self.p[0], h); self.num = ""

    def label(self, text):
        """its own clause number; a lettered item takes the number of the clause above it: (a) under 1.11 -> 1.11(a)."""
        lab = label_of(text)
        if re.fullmatch(r"\d+(\.\d+)*|(Section|SECTION|Clause|CLAUSE) \S+", lab): self.num = lab
        elif lab.startswith("(") and self.num: return self.num + lab
        return lab


# ---------- .docx (tracked changes kept) ----------
def _runs(p):
    """(state, text, bold, author, date) for each run of a paragraph, in order; state = normal | ins | del."""
    out = []
    for r in p.iter(W + "r"):
        state, author, date, a = "normal", "", "", r.getparent()
        while a is not None and a is not p:
            if a.tag in (W + "ins", W + "moveTo"): state = "ins"
            elif a.tag in (W + "del", W + "moveFrom"): state = "del"
            if state != "normal":
                author, date = a.get(W + "author", ""), a.get(W + "date", "")[:10]; break
            a = a.getparent()
        s = []
        for c in r:
            if c.tag in (W + "t", W + "delText"): s.append(c.text or "")
            elif c.tag in (W + "tab", W + "br", W + "cr"): s.append(" ")
            elif c.tag == W + "noBreakHyphen": s.append("-")
        b = r.find(W + "rPr/" + W + "b")
        bold = b is not None and b.get(W + "val", "true") not in ("0", "false")
        if s: out.append((state, "".join(s), bold, author, date))
    return out


def _views(runs):
    """current text, original text, marked text, changes."""
    cur, old, mk, ch = [], [], [], []
    for st, t, _, au, da in runs:
        if st != "del": cur.append(t)
        if st != "ins": old.append(t)
        if st == "normal": mk.append(t)
        else:
            if ch and ch[-1][0] == st and ch[-1][2] == au and mk and mk[-1].endswith(("-]", "+}")):  # merge adjacent runs of one change
                ch[-1] = (st, ch[-1][1] + t, au, da); mk[-1] = mk[-1][:-2] + t + mk[-1][-2:]
            else:
                ch.append((st, t, au, da)); mk.append(f"[-{t}-]" if st == "del" else f"{{+{t}+}}")
    return _clean("".join(cur)), _clean("".join(old)), _clean("".join(mk)), [(s, _clean(t), a, d) for s, t, a, d in ch if _clean(t)]


def _is_heading(p, runs, text):
    st = p.find(W + "pPr/" + W + "pStyle")
    if st is not None and re.match(r"(?i)(heading|title)", st.get(W + "val", "")): return True
    if not text or len(text) > 120 or text.endswith((",", ";")): return False
    vis = [r for r in runs if r[0] != "del" and r[1].strip()]
    if vis and all(r[2] for r in vis) and len(text.split()) <= 14: return True
    return text.isupper() and len(text) >= 4 and len(text.split()) <= 14


def _blocks(el):
    for c in el:
        if c.tag in (W + "p", W + "tbl"): yield c
        elif c.tag in (W + "sdt", W + "sdtContent", W + "customXml"): yield from _blocks(c)


def parse_docx(path, name):
    x = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    units, path_ = [], _Path()
    for b in _blocks(x.find(W + "body")):
        if b.tag == W + "p":
            runs = _runs(b); text, before, marked, ch = _views(runs)
            if not text and not before: continue
            head = _is_heading(b, runs, text or before)
            if head: path_.see(text or before)
            units.append(Unit(name, len(units), "p", text, before, marked, ch, head, path_.label(text or before), path_.p))
        else:  # a table: one unit per row; a header row's cells name the cells below
            header = None
            for tr in b.iter(W + "tr"):
                cells = [_runs(tc) for tc in tr.findall(W + "tc")]
                views = [_views(r) for r in cells]
                if not any(v[0] or v[1] for v in views): continue
                if header is None and all(v[0] for v in views) and all(len(v[0]) < 60 for v in views) and len(views) > 1:
                    header = [v[0] for v in views]
                def join(k):
                    vals = [v[k] for v in views]
                    if header and len(header) == len(vals) and vals != header: return " | ".join(f"{h}: {v}" for h, v in zip(header, vals) if v)
                    return " | ".join(v for v in vals if v)
                ch = [c for v in views for c in v[3]]
                units.append(Unit(name, len(units), "row", join(0), join(1), join(2), ch, False, "", path_.p))
    return units


# ---------- .xlsx ----------
def _cell(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime("%Y-%m-%d")
    if isinstance(v, float): return f"{v:,.4f}".rstrip("0").rstrip(".")
    if isinstance(v, int) and not isinstance(v, bool) and abs(v) >= 10000: return f"{v:,}"
    return _clean(str(v))


def parse_xlsx(path, name):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True); units = []
    for ws in wb.worksheets:
        header = None
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [_cell(v) if v is not None else "" for v in row]
            if not any(vals): continue
            if header is None and sum(bool(v) for v in vals) >= 2 and all(not re.fullmatch(r"[\d,.\-$%]+", v) for v in vals if v):
                header = vals; units.append(Unit(name, len(units), "row", " | ".join(v for v in vals if v), label=f"row {r}", path=(f"sheet {ws.title}",))); continue
            if header and len(header) >= len([v for v in vals]):
                t = " | ".join(f"{h}: {v}" if h else v for h, v in zip(header, vals) if v)
            else:
                t = " | ".join(v for v in vals if v)
            units.append(Unit(name, len(units), "row", t, label=f"row {r}", path=(f"sheet {ws.title}",)))
    return units


# ---------- .pptx ----------
def parse_pptx(path, name):
    from pptx import Presentation
    units = []
    for n, s in enumerate(Presentation(path).slides, 1):
        title = _clean(s.shapes.title.text_frame.text) if s.shapes.title is not None and s.shapes.title.has_text_frame else ""
        for sh in s.shapes:
            paras = []
            if sh.has_text_frame: paras = [_clean(p.text if hasattr(p, "text") else "") for p in sh.text_frame.paragraphs]
            elif getattr(sh, "has_table", False) and sh.has_table: paras = [" | ".join(_clean(c.text) for c in row.cells) for row in sh.table.rows]
            for t in paras:
                if t and t != title: units.append(Unit(name, len(units), "slide", t, label=f"slide {n}", path=(f"slide {n}: {title}"[:100],)))
    return units


# ---------- .eml ----------
def _paras(text):
    out = []
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines: continue
        if all(re.match(r"^([-*•]|\d+[.)]|\([a-z0-9]+\))\s", l) for l in lines[1:]) and len(lines) > 2:  # a list: one unit per item
            out.extend(lines)
        else:
            out.append(" ".join(lines))
    return out


def parse_eml(path, name):
    m = email.message_from_bytes(open(path, "rb").read(), policy=email.policy.default)
    att = [p.get_filename() for p in m.iter_attachments() if p.get_filename()]
    head = "; ".join(f"{k}: {_clean(str(m[k]))}" for k in ("From", "To", "Cc", "Date", "Subject") if m[k])
    if att: head += "; Attachments: " + ", ".join(att)
    subj = _clean(str(m["Subject"] or ""))[:100]
    units = [Unit(name, 0, "head", head, path=(f"email: {subj}",))]
    body = m.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body is not None else ""
    if body is not None and body.get_content_type() == "text/html":
        text = html.unescape(re.sub(r"<[^>]+>", "\n", re.sub(r"(?i)<br\s*/?>|</p>", "\n\n", text)))
    for t in _paras(text):
        units.append(Unit(name, len(units), "p", _clean(t), label=label_of(t), path=(f"email: {subj}",)))
    return units


def parse_text(path, name):
    units, path_ = [], _Path()
    for t in _paras(open(path, encoding="utf-8", errors="replace").read()):
        md = t.startswith("#"); t = _clean(t.lstrip("#"))
        head = len(t) < 100 and (md or t.isupper())
        if head: path_.see(t)
        units.append(Unit(name, len(units), "p", t, heading=head, label=path_.label(t), path=path_.p))
    return units


def parse_pdf(path, name):
    if not shutil.which("pdftotext"): return []
    out = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, timeout=120).stdout
    return [Unit(name, i, "p", _clean(t), label=label_of(t)) for i, t in enumerate(_paras(out))]


PARSERS = {".docx": parse_docx, ".xlsx": parse_xlsx, ".xlsm": parse_xlsx, ".pptx": parse_pptx, ".eml": parse_eml,
           ".txt": parse_text, ".md": parse_text, ".pdf": parse_pdf}


def parse(path):
    """The units of one file ([] when the format is unknown); raises on a corrupt file."""
    f = PARSERS.get(os.path.splitext(path)[1].lower())
    return f(path, os.path.basename(path)) if f else []


if __name__ == "__main__":
    for u in parse(sys.argv[1]):
        print(f"[{u.i}] {u.kind}{' H' if u.heading else ''} {u.where()}\n    {u.marked or u.text}"[:600])
